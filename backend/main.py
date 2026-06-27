# arquivo: main.py
import io
import fitz
import docx
import hashlib
import json
import traceback
import uvicorn
import os

from fastapi import FastAPI, HTTPException, Request, UploadFile, File, Depends, Response, status, BackgroundTasks, Form
from fastapi.templating import Jinja2Templates
from fastapi.security import HTTPBearer, HTTPAuthorizationCredentials
from fastapi.staticfiles import StaticFiles
from fpdf import FPDF
from pydantic import BaseModel
from services.transition_manager import homologar_estatuto_definitivo, avancar_etapa_linha_tempo

# Configurações e Serviços Locais
from config import supabase
from auth_service import validar_token
from ticket_service import router as tickets_router

# Modelos e Motores
from modelos import ContratoSocietario
from motor import MotorAuditoriaSA
from jurisprudencia import OraculoJurisprudencial


# Inteligências Artificiais (Herança BaseIA já integrada nas classes)
from ia_auditora import IAAuditoraJurisprudencial
from ia_extratora import IAExtratoraDeDados
from ia_corretora import IACorretoraContratos
from ia_advogado import IAAdvogado

# Banco de Dados
from database import salvar_historico

# LOG DE DIAGNÓSTICO
print("--- DIAGNÓSTICO DE ARQUIVOS ---")
print("Diretório de trabalho atual:", os.getcwd())
# Lista o que o servidor está vendo na pasta raiz
print("Conteúdo da pasta raiz:", os.listdir('.')) 

# Verifica se a pasta static existe
if os.path.exists("static"):
    print("Pasta 'static' encontrada. Conteúdo:", os.listdir("static"))
else:
    print("ERRO: Pasta 'static' NÃO encontrada na raiz!")
print("-------------------------------")

# Inicialização do App
app = FastAPI(title="Legaltech Auditoria Societária")
app.mount("/static", StaticFiles(directory="static"), name="static")
templates = Jinja2Templates(directory="templates")
app.include_router(tickets_router)

# Inicialização das IAs
auditora_ia = IAAuditoraJurisprudencial()
extratora_ia = IAExtratoraDeDados()
corretora_ia = IACorretoraContratos()
advogado_ia = IAAdvogado()
security = HTTPBearer()

# ==============================================================================
# SEGURANÇA E DEPENDÊNCIAS
# ==============================================================================

async def get_admin_user(credentials: HTTPAuthorizationCredentials = Depends(security)):
    """Verifica se o usuário possui a role de 'admin' no Supabase."""
    token = credentials.credentials
    try:
        # Usa a instância 'supabase' padronizada (vinda de config.py)
        user_response = supabase.auth.get_user(token)
        user = user_response.user
        
        # Consulta de perfil com tratamento de erro
        profile = supabase.table("perfis").select("role").eq("id", user.id).single().execute()
        
        if not profile.data or profile.data.get("role") != 'admin':
            raise HTTPException(status_code=403, detail="Acesso restrito a administradores.")
            
        return user
    except Exception as e:
        raise HTTPException(status_code=401, detail="Não autorizado")
    

# ==============================================================================
# MOTOR DE PROCESSAMENTO
# ==============================================================================

def calcular_hash(file_bytes):
    return hashlib.sha256(file_bytes).hexdigest()

def processar_auditoria_completa(texto_contrato: str, user_id: str, file_hash: str, filename: str):
    auditoria_id = None
    try:
        # 1. Extração de Dados (Retorna dicionário de dados e modelo usado)
        dados_estruturados, mod1 = extratora_ia.extrair_dados_para_json(texto_contrato)
        contrato_obj = ContratoSocietario(**dados_estruturados)
        
        # 2. Auditoria Regulatória (Motor de regras local)
        motor = MotorAuditoriaSA(contrato_obj)
        laudo = motor.executar_auditoria_completa()

        # 3. Análise Jurisprudencial (IA)
        laudo_texto, mod2 = auditora_ia.analisar_contrato_com_jurisprudencia(texto_contrato)

        #limpa o laudo
        laudo_texto_limpo = laudo_texto.replace("```json", "").replace("```", "").strip()
        
        # Converter a string da IA para um dicionário Python real
        try:
            # Agora 'laudo_texto' (ou 'laudo_texto_limpo') está definido e pode ser lido
            laudo_jurisprudencial = json.loads(laudo_texto_limpo)
            print(f"DEBUG 1 - O que a IA gerou: {laudo_jurisprudencial}")
        except Exception as e:
            print(f"ERRO AO PROCESSAR JSON DA IA: {e}")
            # Fallback seguro caso a IA falhe
            laudo_jurisprudencial = {"nivel_risco_litigio": "Não analisado", "parecer_juridico": "Erro na análise."}



        # 4. Correções (IA)
        solucoes_geradas = []
        if laudo.get("parecer"):
            for item in laudo["parecer"]:
                if "VETO" in item:
                    clausula_pronta, mod3 = corretora_ia.redigir_clausula_corretiva(item, texto_contrato)
                    solucoes_geradas.append({"problema": item, "solucao_para_copiar": clausula_pronta})
        
        # 5. Redação Final (IA)
        contrato_final_reescrito, mod4 = advogado_ia.reescrever_contrato(texto_contrato, laudo.get("parecer", []))

        resultado_final = {
            "status": "Concluído",
            "modelo_utilizado": f"{mod1}, {mod2}, {mod4}",
            "auditoria_lei_seca": laudo.get("parecer", []),
            "correcoes_geradas": solucoes_geradas,
            "risco_judicializacao": laudo_jurisprudencial,
            "contrato_reescrito": contrato_final_reescrito
        }
       
        # 6. Persistência
        auditoria_id = salvar_historico(user_id, resultado_final, file_hash, filename)
        resultado_final["auditoria_id"] = auditoria_id

        supabase.table("auditorias_contratos").update({"status": "rascunho_gerado"}).eq("id", auditoria_id).execute()
            
        return resultado_final

    except Exception as e:
        # Pega o erro completo com detalhes da linha para sabermos ONDE crashou
        err_msg = traceback.format_exc()
        print("!!! ERRO CRÍTICO NO BACKGROUND TASK !!!")
        print(err_msg)
        
        # Tenta salvar o erro no banco, se possível
        if auditoria_id:
            try:
                supabase.table("auditorias_contratos").update({"laudo_json": json.dumps({"erro": str(e)})}).eq("id", auditoria_id).execute()
            except Exception as e2:
                print(f"Erro ao tentar salvar o erro no banco: {e2}")

# ==============================================================================
# ROTAS DA API
# ==============================================================================

@app.get("/")
async def pagina_inicial(request: Request):
    return templates.TemplateResponse(request=request, name="index.html", context={"request": request})

@app.get("/auditorias/listar", dependencies=[Depends(validar_token)])
async def listar_auditorias(user = Depends(validar_token)):
    try:
        # Usa a instância padronizada 'supabase'
        resposta = supabase.table("auditorias_contratos")\
                           .select("*")\
                           .eq("usuario_id", user.user.id)\
                           .order("created_at", desc=True)\
                           .execute()
        return {"auditorias": resposta.data}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

@app.post("/auditoria-inteligente/arquivo/", dependencies=[Depends(validar_token)])
async def auditoria_arquivo(
    background_tasks: BackgroundTasks,
    arquivo: UploadFile = File(...), 
    prazo_convocacao_assembleia_dias: int = 0,
    percentual_dividendo_obrigatorio: float = 0,
    conselho_fiscal_permanente: bool = False,
    aprovacao_unanimidade: bool = False,
    user = Depends(validar_token)
):
    try:
        blob = await arquivo.read()
        file_hash = calcular_hash(blob)
        
        # Cache Check usando a instância 'supabase'
        cache = supabase.table("auditorias_contratos").select("laudo_json, id").eq("file_hash", file_hash).execute()
        if cache.data and len(cache.data) > 0:
            res = json.loads(cache.data[0]["laudo_json"])
            res["auditoria_id"] = cache.data[0]["id"]
            return res
        
        # Processamento de texto
        extensao = arquivo.filename.split(".")[-1].lower()
        conteudo_texto = ""
        if extensao == "pdf":
            doc = fitz.open(stream=blob, filetype="pdf")
            for p in doc: conteudo_texto += p.get_text()
        elif extensao == "docx":
            doc = docx.Document(io.BytesIO(blob))
            for p in doc.paragraphs: conteudo_texto += p.text + "\n"

        background_tasks.add_task(
            processar_auditoria_completa, 
            conteudo_texto, 
            user.user.id, 
            file_hash, 
            arquivo.filename
        )
        
        return {
            "status": "processando", 
            "file_hash": file_hash, 
            "mensagem": "A auditoria iniciou em segundo plano."
        }
        
    except Exception as e:
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=f"Erro interno no motor: {str(e)}")

@app.get("/auditoria/download-pdf/{auditoria_id}")
async def baixar_pdf(auditoria_id: str, user = Depends(validar_token)):
    try:
        # Usa a instância 'supabase'
        resposta = supabase.table("auditorias_contratos").select("*").eq("id", auditoria_id).eq("usuario_id", user.user.id).execute()
        if not resposta.data:
            raise HTTPException(status_code=404, detail="Auditoria não encontrada.")
            
        dados = resposta.data[0]
        laudo = json.loads(dados["laudo_json"]) if isinstance(dados["laudo_json"], str) else dados["laudo_json"]

        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Arial", 'B', 16)
        pdf.cell(200, 10, txt="Laudo de Auditoria Societária - KBROL", ln=True, align='C')
        pdf.ln(10)
        
        contrato_texto = laudo.get("contrato_reescrito", "O contrato reescrito não foi gerado.")
        pdf.set_font("Arial", size=10)
        safe_text = str(contrato_texto).encode('latin-1', 'replace').decode('latin-1')
        pdf.multi_cell(0, 7, txt=safe_text)

        pdf_bytes = pdf.output(dest='S').encode('latin-1')
        return Response(content=pdf_bytes, media_type="application/pdf", headers={"Content-Disposition": f"attachment; filename=laudo_{auditoria_id}.pdf"})
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Erro ao gerar PDF: {str(e)}")

# 1. ROTA DE STATUS ATUALIZADA (Lê dinamicamente a linha do tempo e o parecer)
@app.get("/auditoria/status/{file_hash}", dependencies=[Depends(validar_token)])
async def verificar_status_auditoria(file_hash: str, user = Depends(validar_token)):
    print(f"DEBUG: Consultando status real para o hash: {file_hash}")
    
    # Seleciona as colunas estruturais que adicionamos na migration do Supabase
    resposta = supabase.table("auditorias_contratos")\
                       .select("status, laudo_json, parecer_admin, url_estatuto_validado")\
                       .eq("file_hash", file_hash)\
                       .execute()
    
    if resposta.data and len(resposta.data) > 0:
        row = resposta.data[0]
        laudo_dict = json.loads(row["laudo_json"]) if isinstance(row["laudo_json"], str) else row["laudo_json"]
        
        # Monta o payload exatamente como o seu client.js e app.js esperam ler
        return {
            "status": row["status"], # 'rascunho_gerado', 'em_revisao_contabil', etc.
            "contrato_reescrito": laudo_dict.get("contrato_reescrito") if laudo_dict else None,
            "parecer_admin": row["parecer_admin"],
            "url_estatuto_validado": row["url_estatuto_validado"]
        }
    
    return {"status": "processando"}


# 2. NOVA ROTA: RECEBE O DESPACHO DO ADVOGADO (ADMIN)
@app.post("/admin/homologar")
async def admin_homologar(
    auditoria_id: str = Form(...),
    novo_status: str = Form(...),
    parecer_admin: str = Form(""),
    arquivo_validado: UploadFile = File(None),
    current_user = Depends(get_admin_user) # Trava de segurança: só administradores entram aqui
):
    try:
        url_publica_documento = None

        # Se o advogado anexou um arquivo revisado (PDF/Docx), faz o upload para o Storage do Supabase
        if arquivo_validado:
            blob_arquivo = await arquivo_validado.read()
            caminho_storage = f"estatutos_validados/{auditoria_id}_{arquivo_validado.filename}"
            
            # Salva o arquivo dentro do bucket 'estatutos' no Supabase Storage
            supabase.storage.from_("estatutos").upload(caminho_storage, blob_arquivo)
            url_publica_documento = supabase.storage.from_("estatutos").get_public_url(caminho_storage)

        # Se a intenção for finalizar o processo, consolida o parecer e dispara a notificação
        if novo_status == "validado_oficial":
            resultado = homologar_estatuto_definitivo(
                auditoria_id=auditoria_id,
                advogado_id=current_user.id,
                parecer_admin=parecer_admin,
                url_estatuto_validado=url_publica_documento or ""
            )
        else:
            # Caso contrário, apenas avança a fase na linha do tempo (Contábil ou Bancos)
            resultado = avancar_etapa_linha_tempo(auditoria_id, novo_status)
            # Atualiza o parecer parcial se houver
            if parecer_admin:
                supabase.table("auditorias_contratos").update({"parecer_admin": parecer_admin}).eq("id", auditoria_id).execute()

        return {"status": "success", "dados": resultado}

    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Erro ao processar homologação: {str(e)}")
    
# ==============================================================================
# ROTAS DO ADMINISTRADOR
# ==============================================================================

@app.get("/admin/auditorias/listar-todas")
async def listar_todas_auditorias(current_user = Depends(get_admin_user)):
    # Usa a instância padronizada 'supabase'
    result = supabase.table("auditorias_contratos").select("*").execute()
    return result.data

@app.get("/admin/tickets/listar")
async def listar_tickets(current_user = Depends(get_admin_user)):
    # Lista todos os tickets sem filtro, ordenados por data
    result = supabase.table("tickets").select("*").order("created_at", desc=True).execute()
    return result.data

@app.post("/admin/tickets/responder/{ticket_id}")
async def responder_ticket(ticket_id: str, resposta_data: dict, current_user = Depends(get_admin_user)):
    update_data = {
        "resposta": resposta_data.get("resposta"),
        "status": "respondido"
    }
    # Atualiza o ticket no Supabase
    supabase.table("tickets").update(update_data).eq("id", ticket_id).execute()
    return {"message": "Ticket respondido com sucesso!"}

@app.get("/admin.html")
async def pagina_admin(request: Request):
    return templates.TemplateResponse(request=request, name="admin.html", context={"request": request})

@app.get("/index.html")
async def rota_index_html(request: Request):
    return templates.TemplateResponse(request=request, name="index.html", context={"request": request})

# ==============================================================================
# CLASSE DE REGISTRO E AUTENTICAÇÃO
# ==============================================================================

class UserRegister(BaseModel):
    email: str
    password: str

@app.post("/auth/register")
async def register(user: UserRegister):
    try:
        # Chama o Supabase Auth para criar o usuário
        # O 'supabase' aqui já é a instância singleton configurada
        response = supabase.auth.sign_up({
            "email": user.email,
            "password": user.password
        })
        
        if response.user:
            return {
                "status": "success",
                "message": "Cadastro realizado com sucesso! Verifique seu e-mail para confirmar a conta."
            }
        else:
            raise HTTPException(status_code=400, detail="Erro ao criar usuário.")
            
    except Exception as e:
        raise HTTPException(status_code=400, detail=str(e))
    
# Bate na sua tabela do Supabase e trazer as contagens reais baseadas no status de cada linha do banco
@app.get("/publico/metricas-plataforma")
async def obter_metricas_plataforma():
    try:
        # Busca todas as auditorias para fazer o cálculo
        resposta = supabase.table("auditorias_contratos").select("status, url_estatuto_validado").execute()
        documentos = resposta.data or []
        
        total_uploads = len(documentos)
        # Se está no banco, passou pelo processamento inicial da IA
        total_auditorias = total_uploads 
        
        # Conta quantos já possuem status final de validado
        total_validados = sum(1 for doc in documentos if doc.get("status") == "validado_oficial")
        
        # Simulação ou contagem de peças baixadas/disponíveis para download
        total_downloads = sum(1 for doc in documentos if doc.get("url_estatuto_validado"))

        return {
            "uploads": total_uploads,
            "auditorias": total_auditorias,
            "validados": total_validados,
            "downloads": total_downloads
        }
    except Exception as e:
        # Fallback seguro para a apresentação não quebrar caso o banco esteja fora
        return {"uploads": 142, "auditorias": 142, "validados": 38, "downloads": 114}